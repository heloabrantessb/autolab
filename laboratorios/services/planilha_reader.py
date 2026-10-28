"""Leitura limitada de tabelas, sem Django, limpeza, escrita ou execução de fórmulas."""

import csv
from dataclasses import dataclass
from datetime import date, datetime, time
from io import BytesIO, StringIO
import re
from xml.etree.ElementTree import ParseError
from zipfile import BadZipFile, ZipFile

from defusedxml.common import DefusedXmlException
from defusedxml.ElementTree import iterparse
from openpyxl import load_workbook
from openpyxl.utils.cell import column_index_from_string
from openpyxl.utils.exceptions import InvalidFileException

from .zip_validator import ArquivoInvalido, LimitesZIP, inspecionar_zip


class PlanilhaInvalida(Exception):
    """Entrada fora do perfil de leitura permitido."""


@dataclass(frozen=True)
class LimitesPlanilha:
    upload_bytes: int = 20 * 1024 * 1024
    linhas: int = 10000  # dados; cabeçalho não entra nesta contagem
    colunas: int = 30
    celula_caracteres: int = 32768
    xlsx_entradas: int = 100
    xlsx_total_bytes: int = 50 * 1024 * 1024
    xlsx_entrada_bytes: int = 20 * 1024 * 1024


@dataclass(frozen=True)
class TabelaLida:
    formato: str
    cabecalho: tuple[str, ...]
    linhas: tuple[tuple[str, ...], ...]

    @property
    def previa(self):
        return self.linhas[:10]


def _texto(valor, limites):
    if valor is None:
        return ""
    if not isinstance(valor, (str, int, float, date, datetime, time)):
        raise PlanilhaInvalida("A tabela contém um tipo de célula não suportado neste perfil.")
    texto = valor.isoformat() if isinstance(valor, (date, datetime, time)) else str(valor)
    if len(texto) > limites.celula_caracteres:
        raise PlanilhaInvalida("Uma célula excede o limite de caracteres permitido.")
    if any(ord(c) < 32 and c not in "\t\r\n" for c in texto):
        raise PlanilhaInvalida("A tabela contém caracteres de controle não permitidos.")
    return texto


def _tabela(registros, formato, limites):
    cabecalho = next(registros, ())
    if not cabecalho or len(cabecalho) > limites.colunas:
        raise PlanilhaInvalida(f"Informe um cabeçalho com até {limites.colunas} colunas.")
    cabecalho = tuple(_texto(c, limites) for c in cabecalho)
    nomes = [c.strip().casefold() for c in cabecalho]
    if any(not c for c in nomes) or len(set(nomes)) != len(nomes):
        raise PlanilhaInvalida("O cabeçalho deve ter nomes preenchidos e únicos.")
    linhas = []
    for registro in registros:
        if len(linhas) >= limites.linhas:
            raise PlanilhaInvalida(f"A tabela excede o limite de {limites.linhas} linhas de dados.")
        if not registro:  # linha física vazia de CSV também é preservada
            registro = [""] * len(cabecalho)
        if len(registro) != len(cabecalho):
            raise PlanilhaInvalida("Cada linha deve corresponder às colunas do cabeçalho.")
        linhas.append(tuple(_texto(c, limites) for c in registro))
    if not linhas:
        raise PlanilhaInvalida("Inclua pelo menos uma linha de dados além do cabeçalho.")
    return TabelaLida(formato, cabecalho, tuple(linhas))


def _csv(dados, limites):
    try:
        texto = dados.decode("utf-8-sig")
        # Não alteramos csv.field_size_limit: sua configuração é global ao processo.
        registros = csv.reader(StringIO(texto, newline=""), delimiter=",", strict=True)
        return _tabela(registros, "CSV", limites)
    except UnicodeDecodeError as exc:
        raise PlanilhaInvalida("O CSV deve usar codificação UTF-8, com ou sem BOM.") from exc
    except csv.Error as exc:
        raise PlanilhaInvalida("O CSV está malformado ou contém uma célula muito longa.") from exc


_NS = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
_COORDENADA = re.compile(r"([A-Z]{1,3})([1-9][0-9]{0,6})\Z")
_TIPO_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"


def _xml_limitado(arquivo, limites):
    """Bloqueia XML ativo e dimensões esparsas antes de abrir com openpyxl."""
    linha, coluna, anterior, em_dados, elementos = 0, 0, 0, False, 0
    pilha = []
    for evento, elemento in iterparse(
        arquivo, events=("start", "end"),
        forbid_dtd=True, forbid_entities=True, forbid_external=True,
    ):
        tag = elemento.tag
        if evento == "start":
            pilha.append(elemento)
            if len(pilha) > 128:
                raise PlanilhaInvalida("O XML da planilha excede a profundidade permitida.")
            elementos += 1
            if elementos > 1000000:
                raise PlanilhaInvalida("O XML da planilha excede a complexidade permitida.")
            if elemento.get("TargetMode", "").casefold() == "external":
                raise PlanilhaInvalida("XLSX com vínculos externos não é permitido neste perfil.")
            tipo = elemento.get("ContentType", "")
            if "macroenabled" in tipo.casefold() or "vba" in tipo.casefold():
                raise PlanilhaInvalida("Planilhas com macros não são permitidas.")
            if tipo.endswith(".main+xml") and "spreadsheetml" in tipo and tipo != _TIPO_XLSX:
                raise PlanilhaInvalida("Envie uma pasta de trabalho XLSX, sem modelo ou macros.")
            if tag == _NS + "mergeCell":
                raise PlanilhaInvalida("Desfaça as células mescladas antes de enviar a tabela.")
            if tag == _NS + "sheetData":
                em_dados, anterior = True, 0
            elif em_dados and tag == _NS + "row":
                valor = elemento.get("r", "")
                if not valor.isascii() or not valor.isdigit() or len(valor) > 7:
                    raise PlanilhaInvalida("A posição de uma linha XLSX é inválida.")
                linha, coluna = int(valor), 0
                if linha <= anterior or linha > limites.linhas + 1:
                    raise PlanilhaInvalida("As posições das linhas XLSX são inválidas ou excedem o limite.")
                anterior = linha
            elif em_dados and tag == _NS + "c":
                match = _COORDENADA.fullmatch(elemento.get("r", ""))
                if not match or int(match[2]) != linha:
                    raise PlanilhaInvalida("A posição de uma célula XLSX é inválida.")
                atual = column_index_from_string(match[1])
                if atual <= coluna or atual > limites.colunas:
                    raise PlanilhaInvalida("As colunas XLSX são inválidas ou excedem o limite.")
                coluna = atual
        else:
            if tag == _NS + "sheetData":
                em_dados = False
            elif tag == _NS + "row":
                linha = 0
            pilha.pop()
            if pilha:
                pilha[-1].remove(elemento)
            elemento.clear()


def _xlsx(dados, limites):
    entrada = BytesIO(dados)
    livro = None
    try:
        inspecionar_zip(entrada, LimitesZIP(
            limites.upload_bytes, limites.xlsx_entradas,
            limites.xlsx_total_bytes, limites.xlsx_entrada_bytes,
        ))
        with ZipFile(entrada) as pacote:
            for info in pacote.infolist():
                nome = info.filename.casefold()
                if "vbaproject" in nome or nome.startswith("xl/embeddings/"):
                    raise PlanilhaInvalida("Macros e objetos incorporados não são permitidos.")
                if not info.is_dir() and not nome.endswith((".xml", ".rels")):
                    raise PlanilhaInvalida("Use XLSX somente com tabela, sem mídia ou componentes binários.")
                if nome.endswith((".xml", ".rels")):
                    with pacote.open(info) as arquivo:
                        _xml_limitado(arquivo, limites)
        entrada.seek(0)
        livro = load_workbook(entrada, read_only=True, data_only=False, keep_links=False)
        if len(livro.sheetnames) != 1 or len(livro.worksheets) != 1:
            raise PlanilhaInvalida("Envie um XLSX com uma única aba de tabela.")
        folha = livro.worksheets[0]
        # Não confiar em <dimension>, que pode esconder células ou anunciar milhões de linhas.
        folha.reset_dimensions()

        def registros():
            largura = None
            for celulas in folha.iter_rows(max_col=limites.colunas + 1):
                valores = [c.value for c in celulas]
                if any(c.data_type == "e" for c in celulas):
                    raise PlanilhaInvalida("Corrija as células com erro antes de enviar a tabela.")
                if largura is None:
                    while valores and valores[-1] is None:
                        valores.pop()
                    if any(not isinstance(v, str) for v in valores):
                        raise PlanilhaInvalida("O cabeçalho XLSX deve conter nomes em texto.")
                    largura = len(valores)
                else:
                    if any(v is not None for v in valores[largura:]):
                        raise PlanilhaInvalida("Há dados XLSX além das colunas do cabeçalho.")
                    valores = valores[:largura]
                yield valores
        return _tabela(registros(), "XLSX", limites)
    except (ArquivoInvalido, BadZipFile, InvalidFileException, DefusedXmlException,
            ParseError, KeyError, ValueError, TypeError, IndexError, OverflowError, OSError) as exc:
        raise PlanilhaInvalida("O XLSX está inválido, corrompido ou fora dos limites permitidos.") from exc
    finally:
        if livro is not None:
            livro.close()
        entrada.close()


def ler_planilha(entrada, formato, limites=LimitesPlanilha()):
    """Lê stream binário, preserva os dados e devolve tabela limitada em memória."""
    if formato not in ("csv", "xlsx"):
        raise PlanilhaInvalida("Envie CSV ou XLSX.")
    partes, total = [], 0
    try:
        entrada.seek(0)
        while parte := entrada.read(min(65536, limites.upload_bytes - total + 1)):
            total += len(parte)
            if total > limites.upload_bytes:
                raise PlanilhaInvalida("A planilha excede o limite de bytes permitido.")
            partes.append(parte)
        if not total:
            raise PlanilhaInvalida("A planilha está vazia.")
        dados = b"".join(partes)
        return _csv(dados, limites) if formato == "csv" else _xlsx(dados, limites)
    finally:
        entrada.seek(0)

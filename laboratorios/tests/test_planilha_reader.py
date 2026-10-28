"""Leitura real e limites adversariais de CSV/XLSX, sem banco ou escrita do serviço."""

from dataclasses import replace
from datetime import date
from io import BytesIO
from pathlib import Path
from unittest import TestCase
from unittest.mock import patch
from zipfile import ZIP_DEFLATED, ZipFile

from openpyxl import Workbook

from laboratorios.services.planilha_reader import (
    LimitesPlanilha, PlanilhaInvalida, ler_planilha,
)


def xlsx_bytes(linhas, segunda_aba=False):
    livro = Workbook()
    for linha in linhas:
        livro.active.append(linha)
    if segunda_aba:
        livro.create_sheet("Outra")
    saida = BytesIO()
    livro.save(saida)
    livro.close()
    return saida.getvalue()


def alterar_zip(dados, nome, transformar):
    saida = BytesIO()
    with ZipFile(BytesIO(dados)) as origem, ZipFile(saida, "w", ZIP_DEFLATED) as destino:
        for item in origem.infolist():
            conteudo = origem.read(item)
            destino.writestr(item.filename, transformar(conteudo) if item.filename == nome else conteudo)
    return saida.getvalue()


class TestLeitorPlanilhas(TestCase):
    def ler(self, dados, formato="csv", **limites):
        entrada = BytesIO(dados)
        resultado = ler_planilha(entrada, formato, replace(LimitesPlanilha(), **limites))
        self.assertEqual(entrada.tell(), 0)
        return resultado

    def test_fictitious_csv_preserves_all_original_rows_and_values(self):
        dados = (Path(__file__).resolve().parents[2] / "materials/planilhas/clientes.csv").read_bytes()
        tabela = self.ler(dados)
        self.assertEqual(tabela.cabecalho, ("id", "nome", "email", "setor", "observacao"))
        self.assertEqual(len(tabela.linhas), 8)
        self.assertEqual(tabela.linhas[0], tabela.linhas[1])
        self.assertEqual(tabela.linhas[2][1], " Bruno Exemplo ")
        self.assertEqual(tabela.linhas[3][2], "")
        self.assertEqual(tabela.linhas[4], ("", "", "", "", ""))
        self.assertEqual(tabela.linhas[0][0], "001")

    def test_utf8_bom_quoted_commas_and_multiline_values(self):
        tabela = self.ler('\ufeffid,nota\r\n001,"Ação, exemplo\ncontinuação"\r\n'.encode())
        self.assertEqual(tabela.linhas, (("001", "Ação, exemplo\ncontinuação"),))

    def test_preview_is_capped_at_ten_without_dropping_input(self):
        tabela = self.ler(b"id\n" + b"001\n" * 12)
        self.assertEqual(len(tabela.linhas), 12)
        self.assertEqual(len(tabela.previa), 10)

    def test_empty_physical_csv_line_is_preserved(self):
        self.assertEqual(self.ler(b"a,b\n\n1,2\n").linhas, (("", ""), ("1", "2")))

    def test_csv_rejects_invalid_encoding_quotes_width_headers_and_controls(self):
        for dados in (
            b"id,nome\n001,Jo\xe3o\n", b'id,nome\n001,"aberto\n',
            b"a,b\n1,2,3\n", b"a,b\n1\n", b",b\n1,2\n",
            b"a, A \n1,2\n", b"a\n", b"a\nx\x00y\n", b"\n1\n",
        ):
            with self.subTest(dados=dados), self.assertRaises(PlanilhaInvalida):
                self.ler(dados)

    def test_csv_exact_and_excess_rows_columns_cells_and_upload(self):
        dados = b"a,b\n1,2\n3,4\n"
        self.assertEqual(len(self.ler(dados, linhas=2, colunas=2, upload_bytes=len(dados)).linhas), 2)
        for limites in ({"linhas": 1}, {"colunas": 1}, {"upload_bytes": len(dados) - 1}):
            with self.subTest(limites=limites), self.assertRaises(PlanilhaInvalida):
                self.ler(dados, **limites)
        self.assertEqual(self.ler(b"a\n123\n", celula_caracteres=3).linhas, (("123",),))
        with self.assertRaises(PlanilhaInvalida):
            self.ler(b"a\n1234\n", celula_caracteres=3)

    def test_empty_and_unsupported_format_rejected(self):
        for dados, formato in ((b"", "csv"), (b"a\n1", "xls")):
            with self.subTest(formato=formato), self.assertRaises(PlanilhaInvalida):
                self.ler(dados, formato)

    def test_stream_rewound_after_invalid_input(self):
        entrada = BytesIO(b"a\n1\n2\n")
        with self.assertRaises(PlanilhaInvalida):
            ler_planilha(entrada, "csv", replace(LimitesPlanilha(), linhas=1))
        self.assertEqual(entrada.tell(), 0)

    def test_xlsx_preserves_text_ids_blanks_spaces_formulas_and_dates(self):
        dados = xlsx_bytes([
            ["id", "nota", "data"], ["001", " Ana ", date(2026, 1, 2)],
            [None, None, None], ["003", '=1+1', None],
        ])
        tabela = self.ler(dados, "xlsx")
        self.assertEqual(tabela.linhas[0], ("001", " Ana ", "2026-01-02T00:00:00"))
        self.assertEqual(tabela.linhas[1], ("", "", ""))
        self.assertEqual(tabela.linhas[2], ("003", "=1+1", ""))

    def test_xlsx_does_not_trust_understated_or_inflated_dimension(self):
        dados = xlsx_bytes([["id", "nome"], ["001", "Ana"], ["002", "Bruno"]])
        for dimensao in (b"A1:A1", b"A1:XFD1048576"):
            adulterado = alterar_zip(dados, "xl/worksheets/sheet1.xml", lambda d: d.replace(b'A1:B3', dimensao))
            with self.subTest(dimensao=dimensao):
                tabela = self.ler(adulterado, "xlsx")
                self.assertEqual(len(tabela.linhas), 2)
                self.assertEqual(tabela.linhas[-1], ("002", "Bruno"))

    def test_xlsx_sparse_rows_and_columns_checked_before_library_load(self):
        dados = xlsx_bytes([["id"], ["001"]])
        for antigo, novo in ((b'r="2"', b'r="1048576"'), (b'r="A2"', b'r="XFD2"')):
            adulterado = alterar_zip(dados, "xl/worksheets/sheet1.xml", lambda d: d.replace(antigo, novo))
            with self.subTest(novo=novo), patch('laboratorios.services.planilha_reader.load_workbook') as carregar:
                with self.assertRaises(PlanilhaInvalida):
                    self.ler(adulterado, "xlsx")
                carregar.assert_not_called()

    def test_xlsx_actual_dimensions_exact_and_excess(self):
        dados = xlsx_bytes([["a", "b"], ["1", "2"], ["3", "4"]])
        self.assertEqual(len(self.ler(dados, "xlsx", linhas=2, colunas=2).linhas), 2)
        for limites in ({"linhas": 1}, {"colunas": 1}):
            with self.subTest(limites=limites), self.assertRaises(PlanilhaInvalida):
                self.ler(dados, "xlsx", **limites)

    def test_xlsx_multiple_sheets_and_invalid_headers_rejected(self):
        entradas = (
            xlsx_bytes([["id"], ["001"]], segunda_aba=True),
            xlsx_bytes([[1], ["001"]]), xlsx_bytes([["id", " ID "], [1, 2]]),
            xlsx_bytes([["id"], ["001", "sem coluna"]]), xlsx_bytes([["id"]]),
        )
        for dados in entradas:
            with self.subTest(tamanho=len(dados)), self.assertRaises(PlanilhaInvalida):
                self.ler(dados, "xlsx")

    def test_xlsx_merged_cells_rejected(self):
        dados = xlsx_bytes([["id"], ["001"]])
        dados = alterar_zip(dados, "xl/worksheets/sheet1.xml", lambda d: d.replace(b'</worksheet>', b'<mergeCells><mergeCell ref="A1:B1"/></mergeCells></worksheet>'))
        with self.assertRaisesRegex(PlanilhaInvalida, "mescladas"):
            self.ler(dados, "xlsx")

    def test_xlsx_dtd_entities_and_external_relationship_rejected_before_load(self):
        dados = xlsx_bytes([["id"], ["001"]])
        entradas = (
            alterar_zip(dados, "xl/worksheets/sheet1.xml", lambda d: b'<!DOCTYPE worksheet [<!ENTITY x "expansao">]>' + d),
            alterar_zip(dados, "_rels/.rels", lambda d: d.replace(b'Type=', b'TargetMode="External" Type=', 1)),
        )
        for conteudo in entradas:
            with patch('laboratorios.services.planilha_reader.load_workbook') as carregar:
                with self.assertRaises(PlanilhaInvalida):
                    self.ler(conteudo, "xlsx")
                carregar.assert_not_called()

    def test_xlsx_macro_content_type_rejected_before_load(self):
        dados = xlsx_bytes([["id"], ["001"]])
        dados = alterar_zip(dados, "[Content_Types].xml", lambda d: d.replace(
            b'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml',
            b'application/vnd.ms-excel.sheet.macroEnabled.main+xml'))
        with patch('laboratorios.services.planilha_reader.load_workbook') as carregar:
            with self.assertRaises(PlanilhaInvalida):
                self.ler(dados, "xlsx")
            carregar.assert_not_called()

    def test_xlsx_zip_limits_enforced_before_load(self):
        dados = xlsx_bytes([["id"], ["001"]])
        for limites in ({"xlsx_entradas": 1}, {"xlsx_total_bytes": 100}, {"xlsx_entrada_bytes": 100}):
            with self.subTest(limites=limites), patch('laboratorios.services.planilha_reader.load_workbook') as carregar:
                with self.assertRaises(PlanilhaInvalida):
                    self.ler(dados, "xlsx", **limites)
                carregar.assert_not_called()

    def test_xlsx_malformed_xml_and_fake_zip_rejected(self):
        dados = xlsx_bytes([["id"], ["001"]])
        entradas = (
            b"apenas texto", alterar_zip(dados, "xl/workbook.xml", lambda d: b"<quebrado"),
            alterar_zip(dados, "[Content_Types].xml", lambda d: b'<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"/>'),
        )
        for conteudo in entradas:
            with self.assertRaises(PlanilhaInvalida):
                self.ler(conteudo, "xlsx")

    def test_xlsx_unknown_binary_component_rejected_before_load(self):
        dados = BytesIO(xlsx_bytes([["id"], ["001"]]))
        with ZipFile(dados, "a", ZIP_DEFLATED) as arquivo:
            arquivo.writestr("xl/oculto.bin", b"nao abrir")
        with patch('laboratorios.services.planilha_reader.load_workbook') as carregar:
            with self.assertRaises(PlanilhaInvalida):
                self.ler(dados.getvalue(), "xlsx")
            carregar.assert_not_called()

    def test_xlsx_error_cells_are_rejected(self):
        dados = xlsx_bytes([["valor"], ["#DIV/0!"]])
        with self.assertRaisesRegex(PlanilhaInvalida, "células com erro"):
            self.ler(dados, "xlsx")

    def test_xlsx_deep_xml_and_repeated_coordinates_rejected_before_load(self):
        dados = xlsx_bytes([["a", "b"], [1, 2]])
        entradas = (
            alterar_zip(dados, "xl/worksheets/sheet1.xml", lambda d: d.replace(b'r="B2"', b'r="A2"')),
            alterar_zip(dados, "docProps/core.xml", lambda d: b"<a>" * 129 + b"</a>" * 129),
        )
        for conteudo in entradas:
            with patch('laboratorios.services.planilha_reader.load_workbook') as carregar:
                with self.assertRaises(PlanilhaInvalida):
                    self.ler(conteudo, "xlsx")
                carregar.assert_not_called()
